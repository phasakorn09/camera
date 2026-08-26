#!/usr/bin/env python3
"""Verify sample Excel data and render a static dashboard preview."""
from __future__ import annotations

import json
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "App_Data" / "test4.xlsx"
CSS = ROOT / "Content" / "css" / "dashboard.css"
PREVIEW = ROOT / "preview.html"


def sheet_rows(ws):
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    header = [str(value).strip() if value is not None else "" for value in rows[0]]
    data = []
    for row in rows[1:]:
        item = {}
        for index, key in enumerate(header):
            item[key] = row[index] if index < len(row) else None
        data.append(item)
    return data


def money(value):
    return f"{value:,.0f} ฿" if float(value).is_integer() else f"{value:,.2f} ฿"


def main():
    wb = load_workbook(XLSX, data_only=True)
    assert set(wb.sheetnames) == {"สินค้า", "ลูกค้า", "พนักงาน", "การขาย"}

    products = sheet_rows(wb["สินค้า"])
    customers = sheet_rows(wb["ลูกค้า"])
    employees = sheet_rows(wb["พนักงาน"])
    sales = sheet_rows(wb["การขาย"])

    assert len(products) == 10, products
    assert len(customers) == 5
    assert len(employees) == 4
    assert len(sales) == 12

    revenue = sum(float(row["รวมเงิน"]) for row in sales)
    stock_value = sum(float(row["ราคาขาย"]) * int(row["จำนวนคงเหลือ"]) for row in products)
    members = sum(1 for row in customers if str(row["สมาชิก"]).lower() == "yes")
    assert abs(revenue - 51635) < 0.01, revenue
    assert abs(stock_value - 311900) < 0.01, stock_value
    assert members == 3

    product_name = {row["ProductID"]: row["ชื่อสินค้า"] for row in products}
    product_category = {row["ProductID"]: row["ประเภท"] for row in products}
    customer_name = {row["CustomerID"]: f"{row['ชื่อ']} {row['นามสกุล']}" for row in customers}

    sales_by_category = defaultdict(float)
    for row in sales:
        sales_by_category[product_category[row["ProductID"]]] += float(row["รวมเงิน"])

    stock_by_category = defaultdict(int)
    for row in products:
        stock_by_category[row["ประเภท"]] += int(row["จำนวนคงเหลือ"])

    low_stock = [row for row in products if int(row["จำนวนคงเหลือ"]) < 50]
    max_cat = max(sales_by_category.values())

    def table(headers, rows):
        head = "".join(f"<th>{h}</th>" for h in headers)
        body = []
        for row in rows:
            cells = []
            for cell in row:
                css = cell[1] if isinstance(cell, tuple) else ""
                value = cell[0] if isinstance(cell, tuple) else cell
                cells.append(f'<td class="{css}">{value}</td>')
            body.append("<tr>" + "".join(cells) + "</tr>")
        return f'<table class="data-table"><thead><tr>{head}</tr></thead><tbody>{"".join(body)}</tbody></table>'

    sales_rows = []
    for row in sorted(sales, key=lambda item: item["วันที่"], reverse=True):
        sales_rows.append(
            [
                row["SaleID"],
                row["วันที่"].strftime("%d/%m/%Y"),
                customer_name[row["CustomerID"]],
                product_name[row["ProductID"]],
                (str(int(row["จำนวน"])), "text-right"),
                (money(float(row["รวมเงิน"])), "text-right amount-positive"),
            ]
        )

    product_rows = []
    for row in products:
        stock = int(row["จำนวนคงเหลือ"])
        stock_css = "text-right stock-low" if stock < 50 else "text-right"
        product_rows.append(
            [
                row["ProductID"],
                row["ชื่อสินค้า"],
                row["ประเภท"],
                (money(float(row["ราคาขาย"])), "text-right"),
                (f"{stock:,}", stock_css),
            ]
        )

    low_rows = [
        [row["ProductID"], row["ชื่อสินค้า"], (str(int(row["จำนวนคงเหลือ"])), "text-right stock-low")]
        for row in sorted(low_stock, key=lambda item: int(item["จำนวนคงเหลือ"]))
    ]
    customer_rows = [
        [
            row["CustomerID"],
            f"{row['ชื่อ']} {row['นามสกุล']}",
            row["เบอร์โทร"],
            (
                '<span class="badge badge-success">สมาชิก</span>'
                if str(row["สมาชิก"]).lower() == "yes"
                else '<span class="badge badge-muted">ทั่วไป</span>'
            ),
        ]
        for row in customers
    ]
    employee_rows = [
        [
            row["EmployeeID"],
            row["ชื่อ"],
            row["ตำแหน่ง"],
            (money(float(row["เงินเดือน"])), "text-right"),
        ]
        for row in employees
    ]

    cat_items = "".join(
        f'''<li class="category-item"><span>{name}</span>
            <div class="category-bar-wrap"><div class="category-bar" style="width: {(value / max_cat) * 100:.1f}%;"></div></div>
            <span class="category-amount">{money(value)}</span></li>'''
        for name, value in sorted(sales_by_category.items(), key=lambda item: item[1], reverse=True)
    )

    category_json = json.dumps(
        {
            "labels": [name for name, _ in sorted(sales_by_category.items(), key=lambda item: item[1], reverse=True)],
            "values": [value for _, value in sorted(sales_by_category.items(), key=lambda item: item[1], reverse=True)],
        },
        ensure_ascii=False,
    )
    stock_json = json.dumps(
        {
            "labels": [name for name, _ in sorted(stock_by_category.items(), key=lambda item: item[1], reverse=True)],
            "values": [value for _, value in sorted(stock_by_category.items(), key=lambda item: item[1], reverse=True)],
        },
        ensure_ascii=False,
    )

    css_href = "Content/css/dashboard.css"
    html = f"""<!DOCTYPE html>
<html lang="th">
<head>
    <meta charset="utf-8" />
    <meta name="viewport" content="width=device-width, initial-scale=1" />
    <title>แดชบอร์ดร้านวัสดุก่อสร้าง | DashboardWebApp</title>
    <link rel="preconnect" href="https://fonts.googleapis.com" />
    <link rel="preconnect" href="https://fonts.gstatic.com" crossorigin />
    <link href="https://fonts.googleapis.com/css2?family=Prompt:wght@400;500;600;700&amp;family=Sarabun:wght@400;500;600;700&amp;display=swap" rel="stylesheet" />
    <link rel="stylesheet" href="{css_href}" />
</head>
<body>
    <div class="app-layout">
        <aside class="sidebar">
            <div class="brand">
                <span class="brand-mark" aria-hidden="true">
                    <svg viewBox="0 0 48 48" width="36" height="36">
                        <rect x="4" y="18" width="40" height="22" rx="3" fill="#f5a524"></rect>
                        <polygon points="4,18 24,6 44,18" fill="#f8d57a"></polygon>
                        <rect x="20" y="26" width="8" height="14" fill="#122033"></rect>
                    </svg>
                </span>
                <div>
                    <strong>BuildBoard</strong>
                    <span>ร้านวัสดุก่อสร้าง</span>
                </div>
            </div>
            <nav class="side-nav">
                <a href="#overview" class="is-active">ภาพรวม</a>
                <a href="#sales">การขาย</a>
                <a href="#inventory">สินค้าคงเหลือ</a>
                <a href="#people">ลูกค้า &amp; พนักงาน</a>
            </nav>
            <div class="sidebar-foot">
                <p>ASP.NET Web Forms</p>
                <p>อ่านไฟล์ .xlsx ด้วย Open XML</p>
            </div>
        </aside>
        <div class="main-column">
            <header class="topbar">
                <div>
                    <p class="eyebrow">Excel Dashboard</p>
                    <h1>แดชบอร์ดร้านวัสดุก่อสร้าง</h1>
                </div>
                <div class="topbar-meta">
                    <span class="status-pill">ออนไลน์</span>
                    <span class="file-chip">{XLSX.name}</span>
                    <span class="time-chip">อัปเดต {datetime.now().strftime("%d %b %Y %H:%M น.")}</span>
                </div>
            </header>
            <section class="toolbar">
                <div class="toolbar-copy">
                    <h2>โหลดข้อมูลจากไฟล์ Excel</h2>
                    <p>เลือกไฟล์ .xlsx ที่มีชีต สินค้า, ลูกค้า, พนักงาน และการขาย หรือใช้ไฟล์ตัวอย่าง test4.xlsx</p>
                </div>
                <div class="toolbar-actions">
                    <label class="file-picker"><span>เลือกไฟล์ .xlsx</span></label>
                    <button class="btn btn-primary" type="button">โหลดข้อมูล</button>
                    <button class="btn btn-ghost" type="button">ใช้ไฟล์ตัวอย่าง</button>
                    <button class="btn btn-ghost" type="button">ดาวน์โหลดตัวอย่าง</button>
                </div>
            </section>
            <section id="overview" class="stats-grid">
                <article class="stat-card blue"><div class="stat-icon">ยอดขาย</div><div class="stat-label">ยอดขายรวม</div><div class="stat-value">{money(revenue)}</div></article>
                <article class="stat-card emerald"><div class="stat-icon">สินค้า</div><div class="stat-label">รายการสินค้า</div><div class="stat-value">{len(products)}</div></article>
                <article class="stat-card amber"><div class="stat-icon">ลูกค้า</div><div class="stat-label">ลูกค้า / สมาชิก</div><div class="stat-value">{len(customers)} / {members}</div></article>
                <article class="stat-card violet"><div class="stat-icon">ทีม</div><div class="stat-label">พนักงาน</div><div class="stat-value">{len(employees)}</div></article>
                <article class="stat-card cyan"><div class="stat-icon">บิล</div><div class="stat-label">รายการขาย</div><div class="stat-value">{len(sales)}</div></article>
                <article class="stat-card rose"><div class="stat-icon">สต็อก</div><div class="stat-label">มูลค่าสต็อก</div><div class="stat-value">{money(stock_value)}</div></article>
            </section>
            <section class="content-grid charts-grid">
                <article class="panel">
                    <div class="panel-header"><h2>ยอดขายตามประเภทสินค้า</h2></div>
                    <div class="panel-body chart-body">
                        <canvas id="categoryChart" height="220"></canvas>
                        <ul class="category-list">{cat_items}</ul>
                    </div>
                </article>
                <article class="panel">
                    <div class="panel-header"><h2>สัดส่วนสต็อกตามประเภท</h2></div>
                    <div class="panel-body chart-body chart-body-donut">
                        <canvas id="stockChart" height="220"></canvas>
                    </div>
                </article>
            </section>
            <section id="sales" class="content-grid">
                <article class="panel">
                    <div class="panel-header"><h2>รายการขายล่าสุด</h2></div>
                    <div class="panel-body">{table(["รหัส", "วันที่", "ลูกค้า", "สินค้า", "จำนวน", "รวมเงิน"], sales_rows)}</div>
                </article>
                <article class="panel">
                    <div class="panel-header"><h2>สินค้าใกล้หมด</h2></div>
                    <div class="panel-body">{table(["รหัส", "ชื่อสินค้า", "คงเหลือ"], low_rows)}</div>
                </article>
            </section>
            <section id="inventory" class="content-grid single">
                <article class="panel">
                    <div class="panel-header"><h2>สินค้าคงเหลือ</h2></div>
                    <div class="panel-body">{table(["รหัส", "ชื่อสินค้า", "ประเภท", "ราคา", "คงเหลือ"], product_rows)}</div>
                </article>
            </section>
            <section id="people" class="content-grid">
                <article class="panel">
                    <div class="panel-header"><h2>ลูกค้า</h2></div>
                    <div class="panel-body">{table(["รหัส", "ชื่อ-นามสกุล", "เบอร์โทร", "สมาชิก"], customer_rows)}</div>
                </article>
                <article class="panel">
                    <div class="panel-header"><h2>พนักงาน</h2></div>
                    <div class="panel-body">{table(["รหัส", "ชื่อ", "ตำแหน่ง", "เงินเดือน"], employee_rows)}</div>
                </article>
            </section>
            <footer class="dashboard-footer">DashboardWebApp — ASP.NET Web Forms · หน้าเริ่มต้น index.aspx · Visual Studio 2026</footer>
        </div>
    </div>
    <script type="application/json" id="category-chart-data">{category_json}</script>
    <script type="application/json" id="stock-chart-data">{stock_json}</script>
    <script src="https://cdn.jsdelivr.net/npm/chart.js@4.4.1/dist/chart.umd.min.js"></script>
    <script src="Content/js/dashboard.js"></script>
</body>
</html>
"""
    PREVIEW.write_text(html, encoding="utf-8")
    print(f"verified {XLSX}")
    print(f"revenue={revenue} stock_value={stock_value} members={members} low_stock={len(low_stock)}")
    print(f"wrote {PREVIEW}")


if __name__ == "__main__":
    main()
