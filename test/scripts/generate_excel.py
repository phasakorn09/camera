#!/usr/bin/env python3
"""Generate App_Data/test4.xlsx sample data for DashboardWebApp."""
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "DashboardWebApp" / "App_Data" / "test4.xlsx"


def style_header(ws, columns):
    header_font = Font(name="Calibri", bold=True, color="FFFFFF", size=12)
    header_fill = PatternFill("solid", fgColor="1E3A5F")
    thin = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )
    for col, title in enumerate(columns, start=1):
        cell = ws.cell(1, col, title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin
    ws.row_dimensions[1].height = 24
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}1"
    ws.freeze_panes = "A2"


def autosize(ws, widths):
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width


def add_table(ws, name, ref):
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)


def main():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()

    products_headers = ["ProductID", "ชื่อสินค้า", "ประเภท", "ราคาขาย", "ต้นทุน", "จำนวนคงเหลือ"]
    products = [
        ("P001", "ปูนซีเมนต์ Portland", "ปูน", 185, 140, 120),
        ("P002", "เหล็กเส้น 12 มม.", "เหล็ก", 285, 210, 80),
        ("P003", "อิฐมอญ", "อิฐ", 3.5, 2.2, 5000),
        ("P004", "สีทาบ้าน 5 ลิตร", "สี", 890, 620, 45),
        ("P005", "ท่อ PVC 4 นิ้ว", "ท่อ", 125, 85, 200),
        ("P006", "สายไฟ THW 2.5 sqmm", "ไฟฟ้า", 18, 12, 800),
        ("P007", "กระเบื้องปูพื้น 60x60", "กระเบื้อง", 185, 125, 350),
        ("P008", "ทรายหยาบ (ลบ.ม.)", "วัสดุ", 650, 420, 40),
        ("P009", "หินคลุก (ลบ.ม.)", "วัสดุ", 720, 480, 35),
        ("P010", "ประตูไม้เต็ง", "ประตูหน้าต่าง", 4500, 3200, 12),
    ]

    ws = wb.active
    ws.title = "สินค้า"
    style_header(ws, products_headers)
    for row_index, row in enumerate(products, start=2):
        for col, value in enumerate(row, start=1):
            ws.cell(row_index, col, value)
            if col in (4, 5):
                ws.cell(row_index, col).number_format = '#,##0.00'
    autosize(ws, [14, 28, 16, 14, 12, 16])
    add_table(ws, "Products", f"A1:F{len(products) + 1}")

    customers_headers = ["CustomerID", "ชื่อ", "นามสกุล", "เบอร์โทร", "สมาชิก"]
    customers = [
        ("C001", "สมชาย", "ใจดี", "081-234-5678", "Yes"),
        ("C002", "วิภา", "สุขสันต์", "089-111-2233", "Yes"),
        ("C003", "ประยุทธ", "ก่อสร้าง", "062-555-8899", "No"),
        ("C004", "นภา", "เรือนงาม", "098-765-4321", "Yes"),
        ("C005", "ธนา", "ธุรกิจ", "091-222-3344", "No"),
    ]
    ws = wb.create_sheet("ลูกค้า")
    style_header(ws, customers_headers)
    for row_index, row in enumerate(customers, start=2):
        for col, value in enumerate(row, start=1):
            ws.cell(row_index, col, value)
    autosize(ws, [14, 16, 16, 16, 12])
    add_table(ws, "Customers", f"A1:E{len(customers) + 1}")

    employees_headers = ["EmployeeID", "ชื่อ", "ตำแหน่ง", "เงินเดือน"]
    employees = [
        ("E001", "อนุชา แสงทอง", "ผู้จัดการ", 35000),
        ("E002", "มณีรัตน์ ใจดี", "พนักงานขาย", 18000),
        ("E003", "ชัยวัฒน์ มีสุข", "พนักงานคลัง", 16000),
        ("E004", "สุดา รักงาน", "แคชเชียร์", 15000),
    ]
    ws = wb.create_sheet("พนักงาน")
    style_header(ws, employees_headers)
    for row_index, row in enumerate(employees, start=2):
        for col, value in enumerate(row, start=1):
            ws.cell(row_index, col, value)
            if col == 4:
                ws.cell(row_index, col).number_format = '#,##0'
    autosize(ws, [14, 22, 18, 14])
    add_table(ws, "Employees", f"A1:D{len(employees) + 1}")

    sales_headers = [
        "SaleID",
        "วันที่",
        "CustomerID",
        "EmployeeID",
        "ProductID",
        "จำนวน",
        "ราคา",
        "รวมเงิน",
    ]
    sales = [
        ("S001", date(2026, 8, 1), "C001", "E002", "P001", 20, 185, 3700),
        ("S002", date(2026, 8, 3), "C002", "E002", "P002", 15, 285, 4275),
        ("S003", date(2026, 8, 5), "C003", "E001", "P003", 500, 3.5, 1750),
        ("S004", date(2026, 8, 7), "C001", "E004", "P004", 4, 890, 3560),
        ("S005", date(2026, 8, 10), "C004", "E002", "P005", 30, 125, 3750),
        ("S006", date(2026, 8, 12), "C005", "E001", "P007", 40, 185, 7400),
        ("S007", date(2026, 8, 15), "C002", "E004", "P006", 100, 18, 1800),
        ("S008", date(2026, 8, 18), "C003", "E003", "P008", 8, 650, 5200),
        ("S009", date(2026, 8, 20), "C004", "E002", "P010", 1, 4500, 4500),
        ("S010", date(2026, 8, 22), "C001", "E004", "P009", 5, 720, 3600),
        ("S011", date(2026, 8, 24), "C005", "E001", "P002", 10, 285, 2850),
        ("S012", date(2026, 8, 25), "C002", "E002", "P001", 50, 185, 9250),
    ]
    ws = wb.create_sheet("การขาย")
    style_header(ws, sales_headers)
    for row_index, row in enumerate(sales, start=2):
        for col, value in enumerate(row, start=1):
            ws.cell(row_index, col, value)
            if col == 2:
                ws.cell(row_index, col).number_format = "YYYY-MM-DD"
            if col in (7, 8):
                ws.cell(row_index, col).number_format = '#,##0.00'
    autosize(ws, [12, 14, 14, 14, 14, 10, 12, 14])
    add_table(ws, "Sales", f"A1:H{len(sales) + 1}")

    wb.save(OUTPUT)
    print(f"Wrote {OUTPUT}")


if __name__ == "__main__":
    main()
